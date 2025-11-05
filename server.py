import tempfile
import uuid
from fpdf import FPDF
from docx import Document
from PyPDF2 import PdfReader
import os
from typing import List, Optional
from pydantic import BaseModel
from motor.motor_asyncio import AsyncIOMotorClient
from fastapi.responses import FileResponse
from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware

# OCR imports
import pytesseract
from pdf2image import convert_from_path
from PIL import Image

# -------------------------
# CONFIG (adjust paths if needed)
# -------------------------
# Tesseract (Windows) - keep this if you installed Tesseract at this path
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# Poppler (where pdftoppm.exe and pdfinfo.exe live). Use your exact path.
POPLER_BIN_PATH = r"C:\Users\Windows 10\Downloads\poppler-25.07.0\Library\bin"
# -------------------------

app = FastAPI(title="AI-Powered Quiz Generation API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# MongoDB setup
MONGODB_URL = os.getenv("MONGODB_URL", "mongodb://localhost:27017")
client = AsyncIOMotorClient(MONGODB_URL)
db = client.quiz_app
quizzes_collection = db.quizzes


@app.on_event("startup")
async def startup_event():
    try:
        await client.admin.command('ping')
        print("✅ MongoDB connected successfully")
    except Exception as e:
        print(f"❌ MongoDB connection failed: {e}")
        raise e


# -------------------------
# Models
# -------------------------
class Question(BaseModel):
    type: str
    question: str
    options: Optional[List[str]] = None
    answer: str


class Quiz(BaseModel):
    id: str
    title: str
    questions: List[Question]
    metadata: dict


# -------------------------
# Helper: OCR-based PDF -> text
# -------------------------
def extract_pdf_text_with_ocr(pdf_path: str) -> str:
    """
    Convert PDF pages to images using pdf2image (poppler) then OCR with pytesseract.
    Raises an HTTPException with actionable message if poppler not found.
    """
    print("🔍 Using OCR to extract text (pdf2image -> pytesseract)...")

    try:
        pages = convert_from_path(pdf_path, poppler_path=POPLER_BIN_PATH)
    except Exception as e:
        # Most likely poppler not found or not accessible; give helpful guidance
        msg = (
            "pdf2image / Poppler error. Ensure Poppler is installed and poppler_path is correct. "
            f"Attempted poppler_path={POPLER_BIN_PATH}. Underlying error: {e}"
        )
        print("❌", msg)
        raise HTTPException(status_code=500, detail=msg)

    full_text = []
    for page_image in pages:
        try:
            text = pytesseract.image_to_string(page_image)
            full_text.append(text)
        except Exception as e:
            # If OCR fails for a page, continue with others
            print("⚠️ OCR failed for a page:", e)
            continue
        finally:
            # attempt to explicitly close PIL Image to free memory
            try:
                page_image.close()
            except Exception:
                pass

    combined = "\n".join(full_text).strip()
    return combined


# -------------------------
# Helper: Try PyPDF2 first, fallback to OCR
# -------------------------
def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF using PyPDF2; if that returns empty,
    fallback to OCR via extract_pdf_text_with_ocr.
    """
    print("📄 Extracting PDF:", file_path)
    try:
        reader = PdfReader(file_path)
    except Exception as e:
        print("⚠️ PyPDF2 could not open PDF:", e)
        # try OCR anyway
        return extract_pdf_text_with_ocr(file_path)

    extracted_text = []
    try:
        for page in reader.pages:
            try:
                text = page.extract_text()
                if text:
                    extracted_text.append(text)
            except Exception:
                continue
    except Exception as e:
        print("⚠️ Error iterating PDF pages:", e)

    joined = "\n".join(extracted_text).strip()
    if joined:
        print("✅ Extracted text using PyPDF2")
        return joined

    # fallback to OCR
    print("⚠️ PyPDF2 returned empty — switching to OCR")
    return extract_pdf_text_with_ocr(file_path)


# -------------------------
# DOCX extractor
# -------------------------
def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = Document(file_path)
    except Exception as e:
        print("❌ Failed to open DOCX:", e)
        raise HTTPException(
            status_code=400, detail="Unable to read .docx file")

    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(paragraphs).strip()


# -------------------------
# Simple Quiz generator
# -------------------------
def generate_quiz_simple(text: str, quiz_type: str, num_questions: int):
    sentences = [s.strip()
                 for s in text.replace("\n", " ").split(".") if s.strip()]

    if len(sentences) == 0:
        raise HTTPException(status_code=400, detail="Extracted text is empty.")

    questions = []
    for i in range(min(num_questions, len(sentences))):
        sentence = sentences[i]

        if quiz_type == "mcq":
            questions.append({
                "type": "mcq",
                "question": f"What does this sentence describe?\n\"{sentence}\"",
                "options": [
                    "A) " + sentence,
                    "B) Incorrect Option 1",
                    "C) Incorrect Option 2",
                    "D) Incorrect Option 3",
                ],
                "answer": sentence
            })

        elif quiz_type == "true_false":
            questions.append({
                "type": "true_false",
                "question": f"Is this statement true?\n\"{sentence}\"",
                "answer": "True"
            })

        elif quiz_type == "fill_blank":
            parts = sentence.split()
            first_word = parts[0] if parts else ""
            blank_sentence = sentence.replace(
                first_word, "_____") if first_word else sentence
            questions.append({
                "type": "fill_blank",
                "question": blank_sentence,
                "answer": first_word
            })

        elif quiz_type == "identification":
            questions.append({
                "type": "identification",
                "question": f"What is being talked about here?\n\"{sentence}\"",
                "answer": sentence
            })

    return questions[:num_questions]


# -------------------------
# PDF output generator (unchanged)
# -------------------------
def generate_pdf(quiz: Quiz) -> str:
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    pdf.cell(200, 10, txt=f"Quiz: {quiz.title}", ln=True, align='C')
    pdf.ln(10)

    for i, q in enumerate(quiz.questions, 1):
        pdf.multi_cell(0, 10, txt=f"{i}. {q.question}", align="L")
        if q.options:
            for opt in q.options:
                pdf.multi_cell(0, 10, txt=f"   {opt}", align="L")
        pdf.multi_cell(0, 10, txt=f"Answer: {q.answer}", align="L")
        pdf.ln(3)

    pdf.output(temp_file.name)
    return temp_file.name


# -------------------------
# API Endpoints
# -------------------------

@app.post("/api/quizzes/generate")
async def generate_quiz(
    file: Optional[UploadFile] = File(None),
    text: Optional[str] = Form(None),
    quiz_type: str = Form(...),
    num_questions: int = Form(...)
):
    if not text and not file:
        raise HTTPException(status_code=400, detail="Provide text or file.")

    extracted_text = text or ""

    # If file supplied, save to a real temp dir and extract
    if file:
        # safe filename (basic)
        safe_filename = os.path.basename(file.filename)
        temp_dir = tempfile.gettempdir()
        file_path = os.path.join(temp_dir, safe_filename)

        try:
            with open(file_path, "wb") as f:
                f.write(await file.read())
        except Exception as e:
            print("❌ Failed to save uploaded file:", e)
            raise HTTPException(
                status_code=500, detail="Failed to save uploaded file.")

        # check extension case-insensitively
        ext = os.path.splitext(safe_filename)[1].lower()

        try:
            if ext == ".pdf":
                extracted_text = extract_text_from_pdf(file_path)
            elif ext == ".docx":
                extracted_text = extract_text_from_docx(file_path)
            else:
                # support txt quickly
                if ext == ".txt":
                    try:
                        with open(file_path, "r", encoding="utf-8") as tf:
                            extracted_text = tf.read()
                    except Exception:
                        with open(file_path, "r", encoding="latin-1") as tf:
                            extracted_text = tf.read()
                else:
                    raise HTTPException(
                        status_code=400, detail="Unsupported file type.")
        finally:
            # try to remove the uploaded temp file; ignore errors
            try:
                os.remove(file_path)
            except Exception:
                pass

    extracted_text = (extracted_text or "").strip()
    if not extracted_text:
        raise HTTPException(status_code=400, detail="Extracted text is empty.")

    # generate quiz
    questions = generate_quiz_simple(extracted_text, quiz_type, num_questions)

    quiz_id = str(uuid.uuid4())
    quiz = Quiz(
        id=quiz_id,
        title=f"Generated {quiz_type.capitalize()} Quiz",
        questions=[Question(**q) for q in questions],
        metadata={"source": "file" if file else "text",
                  "num_questions": num_questions}
    )

    await quizzes_collection.insert_one(quiz.dict())

    return {"quiz_id": quiz_id, "message": "Quiz generated successfully"}


@app.get("/api/quizzes/{quiz_id}")
async def get_quiz(quiz_id: str):
    quiz_data = await quizzes_collection.find_one({"id": quiz_id})
    if not quiz_data:
        raise HTTPException(status_code=404, detail="Quiz not found")
    return Quiz(**quiz_data)


@app.get("/api/quizzes/{quiz_id}/export")
async def export_quiz(quiz_id: str):
    quiz_data = await quizzes_collection.find_one({"id": quiz_id})
    if not quiz_data:
        raise HTTPException(status_code=404, detail="Quiz not found")

    quiz = Quiz(**quiz_data)
    pdf_path = generate_pdf(quiz)
    return FileResponse(pdf_path, media_type="application/pdf", filename=f"quiz_{quiz_id}.pdf")
